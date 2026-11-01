
def create_deck(title_ch=None, title_lesson=None, lesson_num=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

# -*- coding: utf-8 -*-
"""
generate_teaching_slides.py
Generates 7 complete teaching presentation decks (.pptx) adhering to:
- Official Yukunthor High School NGS branding (Preah Yukunthor logo, teacher profile photo, Mr. Nuon Bonorng)
- Creative visual styling inspired by the template (16:9 widescreen, cards, badges, formula boxes)
- NO E-School logo or mentions
- Pedagogy structure: Objectives -> Recall -> Concepts -> Worked Examples -> MCQs -> Student Worksheets -> Summary -> Homework
- Generates js/teaching-slides-data.js for the interactive Web Dashboard viewer
"""

import os, sys, json
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# Palette
NAVY_DARK = RGBColor(15, 42, 61)       # #0F2A3D
TEAL_DEEP = RGBColor(15, 81, 116)      # #0F5174
GREEN_EMERALD = RGBColor(46, 125, 79)  # #2E7D4F
RED_CRIMSON = RGBColor(192, 39, 45)    # #C0272D
AMBER_GOLD = RGBColor(217, 119, 6)     # #D97706
BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC
CARD_BG = RGBColor(241, 245, 249)      # #F1F5F9
CARD_BORDER = RGBColor(203, 213, 225)  # #CBD5E1
TEXT_DARK = RGBColor(30, 41, 59)       # #1E293B
TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
WHITE = RGBColor(255, 255, 255)
BLUE_ACCENT = RGBColor(37, 99, 235)    # #2563EB
PURPLE_CARD = RGBColor(124, 58, 237)   # #7C3AED

FONT_MUOL = "Khmer OS Muol Light"
FONT_BODY = "Khmer OS Siemreap"
FONT_MATH = "Cambria Math"

LOGO_PATH = "images/yukunthor-logo.jpg"
PROFILE_PATH = "images/teacher-profile.jpg"

def add_header(slide, ch_title, lesson_title):
    # Top banner
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY_DARK
    top_bar.line.fill.background()
    
    # Accent gold line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.1), Inches(13.333), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = AMBER_GOLD
    accent.line.fill.background()
    
    # School logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(0.3), Inches(0.12), height=Inches(0.86))
    
    # Header text
    tb = slide.shapes.add_textbox(Inches(1.3), Inches(0.1), Inches(9.0), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "វិទ្យាល័យព្រះយុគន្ធរ • សាលារៀនជំនាន់ថ្មី (NGS) • គណិតវិទ្យា ថ្នាក់ទី ១២"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(11)
    p1.font.color.rgb = RGBColor(186, 230, 253)
    
    p2 = tf.add_paragraph()
    p2.text = f"{ch_title} — {lesson_title}"
    p2.font.name = FONT_MUOL
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

    # Teacher tag
    tb_t = slide.shapes.add_textbox(Inches(10.3), Inches(0.15), Inches(2.7), Inches(0.8))
    tf_t = tb_t.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "បង្រៀនដោយ៖"
    p_t.font.name = FONT_BODY
    p_t.font.size = Pt(10)
    p_t.font.color.rgb = RGBColor(203, 213, 225)
    p_t.alignment = PP_ALIGN.RIGHT
    p_t2 = tf_t.add_paragraph()
    p_t2.text = "លោកគ្រូ នួន បូណង"
    p_t2.font.name = FONT_MUOL
    p_t2.font.size = Pt(12)
    p_t2.font.color.rgb = AMBER_GOLD
    p_t2.alignment = PP_ALIGN.RIGHT

def add_slide_title(slide, title, category=None):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if category:
        run_cat = p.add_run()
        run_cat.text = f"[{category}] "
        run_cat.font.name = FONT_BODY
        run_cat.font.size = Pt(13)
        run_cat.font.bold = True
        run_cat.font.color.rgb = TEAL_DEEP
    run_t = p.add_run()
    run_t.text = title
    run_t.font.name = FONT_MUOL
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = NAVY_DARK

def create_cover_slide(prs, ch_title, lesson_title, lesson_subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Background full navy
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DARK
    bg.line.fill.background()
    
    # Left banner container (white card)
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(3.8), Inches(5.9))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = AMBER_GOLD
    left_card.line.width = Pt(2.5)
    
    # Teacher Profile photo
    if os.path.exists(PROFILE_PATH):
        slide.shapes.add_picture(PROFILE_PATH, Inches(1.2), Inches(1.2), width=Inches(3.0))
    
    # Teacher label box
    tb_name = slide.shapes.add_textbox(Inches(0.9), Inches(4.5), Inches(3.6), Inches(2.0))
    tf_n = tb_name.text_frame
    tf_n.word_wrap = True
    
    pn1 = tf_n.paragraphs[0]
    pn1.text = "លោកគ្រូ នួន បូណង"
    pn1.font.name = FONT_MUOL
    pn1.font.size = Pt(16)
    pn1.font.bold = True
    pn1.font.color.rgb = NAVY_DARK
    pn1.alignment = PP_ALIGN.CENTER
    
    pn2 = tf_n.add_paragraph()
    pn2.text = "គ្រូបង្រៀនគណិតវិទ្យា ថ្នាក់ទី ១២\nវិទ្យាល័យព្រះយុគន្ធរ\nសាលារៀនជំនាន់ថ្មី (NGS)"
    pn2.font.name = FONT_BODY
    pn2.font.size = Pt(11)
    pn2.font.color.rgb = TEXT_MUTED
    pn2.alignment = PP_ALIGN.CENTER
    
    # Right side: School Branding & Lesson Title
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(5.0), Inches(0.9), height=Inches(1.2))
        
    tb_header = slide.shapes.add_textbox(Inches(6.4), Inches(0.9), Inches(6.2), Inches(1.2))
    tf_h = tb_header.text_frame
    tf_h.word_wrap = True
    ph1 = tf_h.paragraphs[0]
    ph1.text = "ព្រះរាជាណាចក្រកម្ពុជា • ជាតិ សាសនា ព្រះមហាក្សត្រ\nក្រសួងអប់រំ យុវជន និងកីឡា"
    ph1.font.name = FONT_MUOL
    ph1.font.size = Pt(11)
    ph1.font.color.rgb = RGBColor(226, 232, 240)
    
    ph2 = tf_h.add_paragraph()
    ph2.text = "វិទ្យាល័យព្រះយុគន្ធរ • សាលារៀនជំនាន់ថ្មី (NGS)"
    ph2.font.name = FONT_MUOL
    ph2.font.size = Pt(13)
    ph2.font.color.rgb = AMBER_GOLD
    
    # Lesson Title Box
    title_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0), Inches(2.4), Inches(7.5), Inches(3.2))
    title_card.fill.solid()
    title_card.fill.fore_color.rgb = RGBColor(20, 55, 80)
    title_card.line.color.rgb = RGBColor(56, 189, 248)
    title_card.line.width = Pt(1.5)
    
    tb_title = slide.shapes.add_textbox(Inches(5.3), Inches(2.6), Inches(6.9), Inches(2.8))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    
    pt1 = tf_t.paragraphs[0]
    pt1.text = ch_title
    pt1.font.name = FONT_MUOL
    pt1.font.size = Pt(16)
    pt1.font.color.rgb = RGBColor(147, 197, 253)
    
    pt2 = tf_t.add_paragraph()
    pt2.text = lesson_title
    pt2.font.name = FONT_MUOL
    pt2.font.size = Pt(24)
    pt2.font.bold = True
    pt2.font.color.rgb = WHITE
    
    if lesson_subtitle:
        pt3 = tf_t.add_paragraph()
        pt3.text = lesson_subtitle
        pt3.font.name = FONT_BODY
        pt3.font.size = Pt(14)
        pt3.font.color.rgb = RGBColor(203, 213, 225)

    # Bottom pill badges
    tb_foot = slide.shapes.add_textbox(Inches(5.0), Inches(5.8), Inches(7.5), Inches(0.8))
    tf_f = tb_foot.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "✨ កិច្ចតែងការ & ស្លាយបង្រៀនស្ដង់ដារ NGS • ឆ្នាំសិក្សា ២០២៦-២០២៧ • រៀបរៀងដោយលោកគ្រូ នួន បូណង"
    pf.font.name = FONT_BODY
    pf.font.size = Pt(11)
    pf.font.color.rgb = RGBColor(148, 163, 184)
    
    return slide

def create_objectives_slide(prs, ch_title, lesson_title, obj_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, "វត្ថុបំណងនៃមេរៀន (Lesson Objectives)", "គោលបំណង")
    
    sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.95), Inches(11.7), Inches(0.4))
    sub.text_frame.paragraphs[0].text = "បន្ទាប់ពីសិក្សាមេរៀននេះចប់ សិស្សានុសិស្សនឹងទទួលបានសមត្ថភាពសំខាន់ៗ ៣ យ៉ាង៖"
    sub.text_frame.paragraphs[0].font.name = FONT_BODY
    sub.text_frame.paragraphs[0].font.size = Pt(13)
    sub.text_frame.paragraphs[0].font.color.rgb = TEXT_MUTED

    cards_info = [
        ("១. ចំណេះដឹង (Knowledge)", obj_data["knowledge"], GREEN_EMERALD, Inches(0.8)),
        ("២. បំណិន (Skills)", obj_data["skills"], TEAL_DEEP, Inches(4.8)),
        ("៣. ឥរិយាបថ (Attitude)", obj_data["attitude"], RED_CRIMSON, Inches(8.8)),
    ]
    
    for title, items, color, x_pos in cards_info:
        # Card body
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(2.45), Inches(3.7), Inches(4.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(2)
        
        # Header banner inside card
        card_header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(2.45), Inches(3.7), Inches(0.7))
        card_header.fill.solid()
        card_header.fill.fore_color.rgb = color
        card_header.line.fill.background()
        
        tb_h = slide.shapes.add_textbox(x_pos, Inches(2.5), Inches(3.7), Inches(0.6))
        p = tb_h.text_frame.paragraphs[0]
        p.text = title
        p.font.name = FONT_MUOL
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        # Content bullet points
        tb_c = slide.shapes.add_textbox(x_pos + Inches(0.2), Inches(3.25), Inches(3.3), Inches(3.4))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        
        for idx, item in enumerate(items):
            p_item = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
            p_item.text = f"• {item}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = TEXT_DARK
            p_item.space_after = Pt(8)

    return slide

def create_recall_slide(prs, ch_title, lesson_title, recall_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, "រំឭកមេរៀនចាស់ & គំនិតគ្រឹះ (Prerequisite Knowledge)", "មូលដ្ឋានគ្រឹះ")
    
    cards = recall_data.get("cards", [])
    num = len(cards)
    width = Inches(11.7 / max(1, num) - 0.2)
    
    for i, c in enumerate(cards):
        x = Inches(0.8 + i * (11.7 / num))
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.4), width, Inches(4.5))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER
        box.line.width = Pt(1.5)
        
        # top bar of card
        c_top = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.4), width, Inches(0.65))
        c_top.fill.solid()
        c_top.fill.fore_color.rgb = TEAL_DEEP
        c_top.line.fill.background()
        
        tb_t = slide.shapes.add_textbox(x, Inches(2.45), width, Inches(0.55))
        p = tb_t.text_frame.paragraphs[0]
        p.text = c["title"]
        p.font.name = FONT_MUOL
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        tb_b = slide.shapes.add_textbox(x + Inches(0.2), Inches(3.15), width - Inches(0.4), Inches(3.6))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        
        for idx, line in enumerate(c["lines"]):
            p_l = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
            p_l.text = line
            p_l.font.name = FONT_MATH if any(sym in line for sym in ['=', '²', '√', 'π', '+', '-']) else FONT_BODY
            p_l.font.size = Pt(12)
            p_l.font.color.rgb = TEXT_DARK
            p_l.space_after = Pt(6)
            
    return slide

def create_concept_slide(prs, ch_title, lesson_title, concept_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, concept_data["title"], concept_data.get("part", "ទ្រឹស្តីបទ & និយមន័យ"))
    
    # Main theory card (Left)
    card_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(6.8), Inches(4.7))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = WHITE
    card_l.line.color.rgb = CARD_BORDER
    card_l.line.width = Pt(1.5)
    
    tb_l = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(6.4), Inches(4.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    
    for idx, item in enumerate(concept_data.get("points", [])):
        p = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p.text = item
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)
        
    # Formula Highlight Card (Right)
    card_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(2.3), Inches(4.6), Inches(4.7))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = RGBColor(238, 242, 255) # soft indigo
    card_r.line.color.rgb = BLUE_ACCENT
    card_r.line.width = Pt(2)
    
    # Header of formula box
    hdr_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(2.3), Inches(4.6), Inches(0.65))
    hdr_r.fill.solid()
    hdr_r.fill.fore_color.rgb = BLUE_ACCENT
    hdr_r.line.fill.background()
    
    tb_hr = slide.shapes.add_textbox(Inches(7.9), Inches(2.35), Inches(4.6), Inches(0.55))
    phr = tb_hr.text_frame.paragraphs[0]
    phr.text = "📌 រូបមន្តគន្លឹះ និងចំណាំសំខាន់"
    phr.font.name = FONT_MUOL
    phr.font.size = Pt(13)
    phr.font.color.rgb = WHITE
    phr.alignment = PP_ALIGN.CENTER
    
    tb_r = slide.shapes.add_textbox(Inches(8.1), Inches(3.1), Inches(4.2), Inches(3.7))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    for idx, f in enumerate(concept_data.get("formulas", [])):
        pf = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
        pf.text = f
        pf.font.name = FONT_MATH
        pf.font.size = Pt(14)
        pf.font.bold = True
        pf.font.color.rgb = NAVY_DARK
        pf.space_after = Pt(10)
        
    return slide

def create_example_slide(prs, ch_title, lesson_title, ex_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, ex_data["title"], "លំហាត់គំរូ")
    
    # Problem statement card (Top)
    prob_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.1))
    prob_card.fill.solid()
    prob_card.fill.fore_color.rgb = RGBColor(240, 249, 255) # light sky
    prob_card.line.color.rgb = TEAL_DEEP
    prob_card.line.width = Pt(1.5)
    
    tb_p = slide.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.3), Inches(0.95))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    pp = tf_p.paragraphs[0]
    run_q = pp.add_run()
    run_q.text = "លំហាត់៖ "
    run_q.font.name = FONT_MUOL
    run_q.font.size = Pt(14)
    run_q.font.color.rgb = TEAL_DEEP
    
    run_text = pp.add_run()
    run_text.text = ex_data["problem"]
    run_text.font.name = FONT_BODY
    run_text.font.size = Pt(13)
    run_text.font.color.rgb = TEXT_DARK
    
    # Solution card (Bottom Left)
    sol_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.45), Inches(8.0), Inches(3.6))
    sol_card.fill.solid()
    sol_card.fill.fore_color.rgb = WHITE
    sol_card.line.color.rgb = CARD_BORDER
    sol_card.line.width = Pt(1.5)
    
    tb_s = slide.shapes.add_textbox(Inches(1.0), Inches(3.55), Inches(7.6), Inches(3.4))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    
    ps_head = tf_s.paragraphs[0]
    ps_head.text = "ដំណោះស្រាយលម្អិត៖"
    ps_head.font.name = FONT_MUOL
    ps_head.font.size = Pt(13)
    ps_head.font.color.rgb = GREEN_EMERALD
    ps_head.space_after = Pt(4)
    
    for step in ex_data.get("solution_steps", []):
        p_step = tf_s.add_paragraph()
        p_step.text = step
        p_step.font.name = FONT_BODY
        p_step.font.size = Pt(12)
        p_step.font.color.rgb = TEXT_DARK
        p_step.space_after = Pt(4)
        
    # Tip card (Bottom Right)
    tip_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(3.45), Inches(3.5), Inches(3.6))
    tip_card.fill.solid()
    tip_card.fill.fore_color.rgb = RGBColor(254, 252, 232) # soft yellow
    tip_card.line.color.rgb = AMBER_GOLD
    tip_card.line.width = Pt(1.5)
    
    tb_tip = slide.shapes.add_textbox(Inches(9.1), Inches(3.55), Inches(3.3), Inches(3.4))
    tf_tip = tb_tip.text_frame
    tf_tip.word_wrap = True
    
    ptip_h = tf_tip.paragraphs[0]
    ptip_h.text = "💡 គន្លឹះដោះស្រាយ"
    ptip_h.font.name = FONT_MUOL
    ptip_h.font.size = Pt(13)
    ptip_h.font.color.rgb = AMBER_GOLD
    ptip_h.space_after = Pt(6)
    
    for tip in ex_data.get("tips", []):
        ptip = tf_tip.add_paragraph()
        ptip.text = f"• {tip}"
        ptip.font.name = FONT_BODY
        ptip.font.size = Pt(11)
        ptip.font.color.rgb = TEXT_DARK
        ptip.space_after = Pt(4)
        
    return slide

def create_mcq_slide(prs, ch_title, lesson_title, mcq_data, mcq_index=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, f"សំណួរពហុជ្រើសរើសទី {mcq_index} (MCQ {mcq_index})", "ពង្រឹងការយល់ដឹង")
    
    # Question Box (Top)
    q_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.3))
    q_card.fill.solid()
    q_card.fill.fore_color.rgb = RGBColor(245, 243, 255) # light purple
    q_card.line.color.rgb = PURPLE_CARD
    q_card.line.width = Pt(2)
    
    tb_q = slide.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.3), Inches(1.2))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    pq = tf_q.paragraphs[0]
    run_lbl = pq.add_run()
    run_lbl.text = f"❓ សំណួរទី {mcq_index} ៖ "
    run_lbl.font.name = FONT_MUOL
    run_lbl.font.size = Pt(14)
    run_lbl.font.color.rgb = PURPLE_CARD
    
    run_qtext = pq.add_run()
    run_qtext.text = mcq_data["question"]
    run_qtext.font.name = FONT_BODY
    run_qtext.font.size = Pt(13)
    run_qtext.font.bold = True
    run_qtext.font.color.rgb = NAVY_DARK
    
    # 4 Choice cards (2x2 grid)
    choices = mcq_data.get("choices", [])
    positions = [
        (Inches(0.8), Inches(3.65), Inches(5.7), Inches(1.2)), # (A)
        (Inches(6.8), Inches(3.65), Inches(5.7), Inches(1.2)), # (B)
        (Inches(0.8), Inches(5.0), Inches(5.7), Inches(1.2)),  # (C)
        (Inches(6.8), Inches(5.0), Inches(5.7), Inches(1.2)),  # (D)
    ]
    
    for i, c in enumerate(choices[:4]):
        x, y, w, h = positions[i]
        is_correct = (i == mcq_data.get("correct_index", 0))
        
        c_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        c_box.fill.solid()
        if is_correct:
            c_box.fill.fore_color.rgb = RGBColor(240, 253, 244) # soft green
            c_box.line.color.rgb = GREEN_EMERALD
            c_box.line.width = Pt(2)
        else:
            c_box.fill.fore_color.rgb = CARD_BG
            c_box.line.color.rgb = CARD_BORDER
            c_box.line.width = Pt(1.5)
            
        tb_c = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), h - Inches(0.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        pc = tf_c.paragraphs[0]
        
        run_opt = pc.add_run()
        prefix = ["(ក)", "(ខ)", "(គ)", "(ឃ)"][i]
        run_opt.text = f"{prefix}  "
        run_opt.font.name = FONT_MUOL
        run_opt.font.size = Pt(13)
        run_opt.font.bold = True
        run_opt.font.color.rgb = GREEN_EMERALD if is_correct else NAVY_DARK
        
        run_c = pc.add_run()
        run_c.text = c
        run_c.font.name = FONT_BODY
        run_c.font.size = Pt(13)
        run_c.font.color.rgb = TEXT_DARK
        
        if is_correct:
            p_badge = tf_c.add_paragraph()
            p_badge.text = f"✓ ចម្លើយត្រឹមត្រូវ! ({mcq_data.get('explanation', '')})"
            p_badge.font.name = FONT_BODY
            p_badge.font.size = Pt(10)
            p_badge.font.bold = True
            p_badge.font.color.rgb = GREEN_EMERALD

    # Bottom note
    tb_b = slide.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.5))
    pb = tb_b.text_frame.paragraphs[0]
    pb.text = "💡 សិស្សលើកដៃឆ្លើយ ឬបោះឆ្នោតតាមទូរស័ព្ទ / ក្តារឆ្នួន រួចគ្រូបកស្រាយមូលហេតុលម្អិត។"
    pb.font.name = FONT_BODY
    pb.font.size = Pt(11)
    pb.font.color.rgb = TEXT_MUTED
    
    return slide

def create_worksheet_slide(prs, ch_title, lesson_title, ws_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, "សន្លឹកកិច្ចការសិស្ស & ការពិភាក្សាជាក្រុម (Student Worksheet)", "អនុវត្តជាក់ស្តែង")
    
    # Instruction bar
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL_DEEP
    bar.line.fill.background()
    
    tb_bar = slide.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.3), Inches(0.6))
    pbar = tb_bar.text_frame.paragraphs[0]
    pbar.text = "⏱️ រយៈពេលពិភាក្សា៖ ១០ នាទី • ពិន្ទុអតិបរមា៖ ១០ ពិន្ទុ • សមាជិកក្នុងក្រុមសហការដោះស្រាយ និងឡើងធ្វើបទបង្ហាញ"
    pbar.font.name = FONT_BODY
    pbar.font.size = Pt(12)
    pbar.font.bold = True
    pbar.font.color.rgb = WHITE
    
    # 3 Group cards
    groups = ws_data.get("groups", [])
    for i, g in enumerate(groups[:3]):
        x = Inches(0.8 + i * 4.0)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.05), Inches(3.7), Inches(3.9))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = [BLUE_ACCENT, GREEN_EMERALD, RED_CRIMSON][i % 3]
        card.line.width = Pt(2)
        
        # Header banner of group card
        g_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.05), Inches(3.7), Inches(0.65))
        g_hdr.fill.solid()
        g_hdr.fill.fore_color.rgb = [BLUE_ACCENT, GREEN_EMERALD, RED_CRIMSON][i % 3]
        g_hdr.line.fill.background()
        
        tb_gh = slide.shapes.add_textbox(x, Inches(3.1), Inches(3.7), Inches(0.55))
        pgh = tb_gh.text_frame.paragraphs[0]
        pgh.text = g["name"]
        pgh.font.name = FONT_MUOL
        pgh.font.size = Pt(13)
        pgh.font.color.rgb = WHITE
        pgh.alignment = PP_ALIGN.CENTER
        
        # Task content
        tb_gc = slide.shapes.add_textbox(x + Inches(0.2), Inches(3.8), Inches(3.3), Inches(3.0))
        tf_gc = tb_gc.text_frame
        tf_gc.word_wrap = True
        
        for idx, task in enumerate(g.get("tasks", [])):
            pt = tf_gc.paragraphs[0] if idx == 0 else tf_gc.add_paragraph()
            pt.text = f"{idx+1}. {task}"
            pt.font.name = FONT_BODY
            pt.font.size = Pt(11)
            pt.font.color.rgb = TEXT_DARK
            pt.space_after = Pt(6)
            
    return slide

def create_summary_slide(prs, ch_title, lesson_title, summary_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, "សេចក្តីសង្ខេប & តារាងរូបមន្តគន្លឹះ (Summary Cheat Sheet)", "សង្ខេបមេរៀន")
    
    rows = summary_data.get("rows", [])
    num_rows = len(rows) + 1
    
    # Create Table
    table_shape = slide.shapes.add_table(num_rows, 4, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.7))
    table = table_shape.table
    table.columns[0].width = Inches(1.0) # No
    table.columns[1].width = Inches(3.5) # Topic
    table.columns[2].width = Inches(4.5) # Formula
    table.columns[3].width = Inches(2.7) # Note
    
    headers = ["ល.រ", "ទម្រង់ / ចំណងជើង", "រូបមន្តគន្លឹះ / លក្ខណៈ", "ចំណាំ / អត្ថន័យ"]
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_MUOL
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
    for r_idx, r in enumerate(rows):
        row_bg = CARD_BG if r_idx % 2 == 0 else WHITE
        for c_idx in range(4):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            p = cell.text_frame.paragraphs[0]
            val = [str(r_idx + 1), r["topic"], r["formula"], r["note"]][c_idx]
            p.text = val
            p.font.name = FONT_MATH if c_idx == 2 else FONT_BODY
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            if c_idx == 0:
                p.alignment = PP_ALIGN.CENTER
                
    return slide

def create_homework_slide(prs, ch_title, lesson_title, hw_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, ch_title, lesson_title)
    add_slide_title(slide, "កិច្ចការផ្ទះ & សូមអរគុណ! (Homework & Closing)", "កិច្ចការបន្ត")
    
    # Homework Card (Left)
    hw_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(6.0), Inches(4.7))
    hw_card.fill.solid()
    hw_card.fill.fore_color.rgb = CARD_BG
    hw_card.line.color.rgb = AMBER_GOLD
    hw_card.line.width = Pt(2)
    
    tb_hw = slide.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(5.4), Inches(4.3))
    tf_hw = tb_hw.text_frame
    tf_hw.word_wrap = True
    
    phw_h = tf_hw.paragraphs[0]
    phw_h.text = "📚 កិច្ចការផ្ទះត្រូវអនុវត្ត"
    phw_h.font.name = FONT_MUOL
    phw_h.font.size = Pt(15)
    phw_h.font.color.rgb = AMBER_GOLD
    phw_h.space_after = Pt(8)
    
    for item in hw_data.get("tasks", []):
        pi = tf_hw.add_paragraph()
        pi.text = f"• {item}"
        pi.font.name = FONT_BODY
        pi.font.size = Pt(13)
        pi.font.color.rgb = TEXT_DARK
        pi.space_after = Pt(6)
        
    p_due = tf_hw.add_paragraph()
    p_due.text = "⏰ កាលបរិច្ឆេទប្រគល់៖ ម៉ោងសិក្សាបន្ទាប់"
    p_due.font.name = FONT_BODY
    p_due.font.size = Pt(12)
    p_due.font.bold = True
    p_due.font.color.rgb = RED_CRIMSON
    p_due.space_before = Pt(8)
    
    # Thank You & Contact Card (Right)
    ty_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.3), Inches(5.3), Inches(4.7))
    ty_card.fill.solid()
    ty_card.fill.fore_color.rgb = NAVY_DARK
    ty_card.line.color.rgb = TEAL_DEEP
    ty_card.line.width = Pt(2)
    
    if os.path.exists(PROFILE_PATH):
        slide.shapes.add_picture(PROFILE_PATH, Inches(8.8), Inches(2.6), height=Inches(1.8))
        
    tb_ty = slide.shapes.add_textbox(Inches(7.4), Inches(4.6), Inches(4.9), Inches(2.2))
    tf_ty = tb_ty.text_frame
    tf_ty.word_wrap = True
    
    pty1 = tf_ty.paragraphs[0]
    pty1.text = "សូមអរគុណសម្រាប់ការយកចិត្តទុកដាក់!"
    pty1.font.name = FONT_MUOL
    pty1.font.size = Pt(14)
    pty1.font.bold = True
    pty1.font.color.rgb = WHITE
    pty1.alignment = PP_ALIGN.CENTER
    
    pty2 = tf_ty.add_paragraph()
    pty2.text = "លោកគ្រូ នួន បូណង • វិទ្យាល័យព្រះយុគន្ធរ NGS\nមុខវិជ្ជា៖ គណិតវិទ្យា ថ្នាក់ទី ១២\nTelegram Channel: @MathNGS_Yukunthor"
    pty2.font.name = FONT_BODY
    pty2.font.size = Pt(11)
    pty2.font.color.rgb = RGBColor(186, 230, 253)
    pty2.alignment = PP_ALIGN.CENTER
    
    return slide

print("Helper functions defined.")
